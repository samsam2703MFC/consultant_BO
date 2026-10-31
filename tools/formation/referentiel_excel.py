#!/usr/bin/env python3
"""Référentiel v25 (Excel) → JSON structuré.

Usage :
    python3 -I tools/formation/referentiel_excel.py atelier_referentiel_v25.xlsx sortie.json

Lit les onglets Formations, Matrice, Compétences, Plans de Formation et Manuel
Opératoire du classeur « L'ATELIER BY — RÉFÉRENTIEL v25 » et écrit :
  postes[]       P1…P12 (code, nom)
  formations[]   FO-xx : nom, priorités business, fréquence, sous-cat., sections,
                 contrôleur, exécutant, procédure MO, attendu par poste (O / R)
  competences[]  C-FAMILLE.Nom : famille, nom, description, FO parent, niveau,
                 postes clefs
  plans[]        PF-xx : nom, logique, format, budget, blocs[] (contenu, check-list,
                 KPIs, FO liées)
  manuel[]       domaines D1…D8 et leurs sous-catégories (priorités, compétences,
                 checklists, FO dérivées)
Rien n'est écrit ailleurs : c'est la source des maquettes et de l'import à venir.
"""
import json
import re
import sys

import openpyxl

PRIORITES = ['TRAFIC', 'RECURRENCE', 'XP', 'FOOD-COST', 'LABOUR-COST', 'OVERHEAD-COST']


def lignes(ws):
    for r in ws.iter_rows(values_only=True):
        if any(c not in (None, '') for c in r):
            yield [('' if c is None else (str(c).strip() if not isinstance(c, float) else c)) for c in r]


def texte(v):
    if isinstance(v, float):
        return str(int(v)) if v == int(v) else str(v)
    return str(v or '').strip()


def puces(v):
    return [x.strip(' •·-\t') for x in texte(v).split('\n') if x.strip(' •·-\t') and x.strip() != '—']


def main(src, dst):
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = {w.title: w for w in wb.worksheets}

    # --- Matrice : postes P1…P12 et attendu O/R par formation
    postes, attendu = [], {}
    for r in lignes(ws['Matrice']):
        if r[0] == '#' and texte(r[1]) == 'ID':
            for c in r[4:]:
                t = texte(c).replace('\n', ' ')
                m = re.match(r'(P\d+)\s*/?\s*(.*)', t)
                if m:
                    postes.append({'code': m.group(1), 'nom': m.group(2).strip()})
            continue
        if postes and re.match(r'FO-\d+', texte(r[1])):
            attendu[texte(r[1])] = {p['code']: texte(r[4 + i]) for i, p in enumerate(postes) if 4 + i < len(r) and texte(r[4 + i]) in ('O', 'R')}

    # --- Formations
    formations = []
    for r in lignes(ws['Formations']):
        if not re.match(r'FO-\d+', texte(r[1])):
            continue
        formations.append({
            'code': texte(r[1]), 'nom': texte(r[2]),
            'priorites': [p for i, p in enumerate(PRIORITES) if r[3 + i] in (1, 1.0, '1', '1.0')],
            'frequence': texte(r[9]), 'sousCategorie': texte(r[10]), 'sections': texte(r[11]),
            'controleur': texte(r[12]), 'executant': texte(r[13]), 'procedure': texte(r[14]),
            'nbEvaluations': int(r[15]) if isinstance(r[15], float) else None,
            'attendu': attendu.get(texte(r[1]), {}),
        })

    # --- Compétences (vue plate, FK FO parent)
    competences = []
    for r in lignes(ws['Compétences']):
        if not re.match(r'C-[A-Z]+\.', texte(r[1])):
            continue
        competences.append({
            'code': texte(r[1]), 'famille': texte(r[2]), 'nom': texte(r[3]), 'description': texte(r[4]),
            'formation': texte(r[5]), 'niveau': texte(r[7]) if len(r) > 7 else '', 'postes': texte(r[8]) if len(r) > 8 else '',
        })

    # --- Plans de formation : un en-tête « ▸ PF-xx — … », sa logique, puis ses blocs
    plans, plan = [], None
    for r in lignes(ws['Plans de Formation']):
        a = texte(r[0])
        m = re.match(r'▸\s*(PF-[A-Z]+)\s*—\s*(.*?)\s*·', a)
        if m:
            plan = {'code': m.group(1), 'nom': m.group(2).strip(), 'logique': '', 'format': '', 'budget': '', 'blocs': []}
            plans.append(plan)
            continue
        if plan is not None and a.startswith('Logique'):
            for part in a.split('\n'):
                for cle, lib in (('logique', 'Logique'), ('format', 'Format'), ('budget', 'Budget')):
                    if part.strip().startswith(lib):
                        plan[cle] = part.split(':', 1)[1].strip() if ':' in part else part.strip()
            continue
        if plan is not None and re.match(r'PF-[A-Z]+', texte(r[1])) and re.match(r'B\d+', texte(r[3])):
            plan['blocs'].append({
                'code': texte(r[3]), 'nom': texte(r[4]), 'contenu': puces(r[5]),
                'checklist': puces(r[6]), 'kpis': puces(r[7]),
                'formations': [texte(x) for x in r[8:11] if re.match(r'FO-\d+', texte(x))],
            })

    # --- Manuel opératoire : domaines et sous-catégories
    manuel, dom = [], None
    for r in lignes(ws['Manuel Opératoire']):
        a = texte(r[0])
        m = re.match(r'▸\s*(D\d+)\s*—\s*(.*?)\s*·', a)
        if m:
            dom = {'code': m.group(1), 'nom': m.group(2).strip(), 'sousCategories': []}
            manuel.append(dom)
            continue
        if dom is not None and re.match(r'D\d+', texte(r[1])) and texte(r[3]):
            dom['sousCategories'].append({
                'nom': texte(r[3]), 'description': texte(r[10]),
                'priorites': [p for i, p in enumerate(PRIORITES) if r[4 + i] in (1, 1.0)],
                'competences': [texte(x) for x in r[11:15] if texte(x)],
                'checklists': [texte(x) for x in r[15:22] if texte(x)],
                'formations': [texte(x) for x in r[22:26] if re.match(r'FO-\d+', texte(x))],
            })

    out = {'source': 'L\'ATELIER BY — RÉFÉRENTIEL v25 (Excel)', 'priorites': PRIORITES, 'postes': postes,
           'formations': formations, 'competences': competences, 'plans': plans, 'manuel': manuel}
    with open(dst, 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
        fh.write('\n')
    print(f"{len(postes)} postes, {len(formations)} formations, {len(competences)} compétences, "
          f"{len(plans)} plans / {sum(len(p['blocs']) for p in plans)} blocs, "
          f"{len(manuel)} domaines / {sum(len(d['sousCategories']) for d in manuel)} sous-cat. -> {dst}")


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
