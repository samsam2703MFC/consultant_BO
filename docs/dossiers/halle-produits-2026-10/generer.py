# La liste des produits de Halle avec leur coefficient : classeur Excel (coloré par groupe, coefficient par zone) et HTML A4 → PDF.
import json, html, datetime, base64
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
SP = '/tmp/claude-0/-home-user-consultant-BO/64c06f7c-f933-5fd1-a986-9b6771ad42a2/scratchpad'
d = json.load(open(SP + '/coef/halle.json'))
LU = datetime.datetime.fromisoformat(d['lu']).strftime('%d/%m/%Y à %H:%M')
S = d['seuils']
def zone(p):   # la règle de vraisemblance du P&L, appliquée ici aussi (les données lues précèdent le déploiement de la règle)
    c = p['coef']
    if c is None: return 'mu'
    if c <= 1 or c > 20: return 'ab'
    return 'ko' if c < S['ko'] else ('att' if c < S['att'] else 'ok')
P = d['produits']
for p in P:
    p['zone'] = zone(p)
    if not p['groupe']: p['groupe'] = p['cat']
P.sort(key=lambda p: (p['groupe'], p['cat'], p['nom']))
ZL = {'ok': 'vert · marge ≥ 60 %', 'att': 'orange · marge 40 à 60 %', 'ko': 'rouge · marge < 40 %', 'mu': 'sans coût de recette ou sans prix', 'ab': 'recette à vérifier'}
ZC = {'ok': ('2D7A3E', 'D7EBD9'), 'att': ('C96A1B', 'FDE6D0'), 'ko': ('C0182B', 'F6D3D6'), 'mu': ('7A736A', 'EEECE8'), 'ab': ('6A3FA0', 'E6DDF2')}
PAL = ['8D1D2C', 'C96A2B', 'D9A441', '8A6A3B', '5B8C3E', '2F7F8F', 'B07F4E', '3B6BB5', '7A7A7A', 'A8407A', 'D25C3D', '4F5B7A', '6E8B3D', '9C5B2E', '3F7D6B', '7B4B94']
def teinte(hexa, f=0.82):
    r, g, b = int(hexa[0:2], 16), int(hexa[2:4], 16), int(hexa[4:6], 16)
    return '%02X%02X%02X' % tuple(int(c + (255 - c) * f) for c in (r, g, b))
groupes = sorted({p['groupe'] for p in P}, key=lambda g: (-sum(1 for p in P if p['groupe'] == g), g))
GC = {g: PAL[i % len(PAL)] for i, g in enumerate(groupes)}
def med(v):
    v = sorted(x for x in v if x is not None)
    if not v: return None
    n = len(v); return v[n // 2] if n % 2 else (v[n // 2 - 1] + v[n // 2]) / 2
def agg(cle):
    out = []
    for k in sorted({p[cle] for p in P}, key=lambda k: (groupes.index(k) if cle == 'groupe' else (groupes.index([p for p in P if p[cle] == k][0]['groupe']), k))):
        ps = [p for p in P if p[cle] == k]; ok = [p['coef'] for p in ps if p['zone'] in ('ok', 'att', 'ko')]
        z = {x: sum(1 for p in ps if p['zone'] == x) for x in ZL}
        out.append(dict(nom=k, groupe=ps[0]['groupe'], n=len(ps), chiffres=len(ok), med=med(ok), mn=min(ok) if ok else None, mx=max(ok) if ok else None, cible=ps[0].get('cible') if cle == 'cat' else None, z=z))
    return out
AC = agg('cat'); AG = agg('groupe')
nOk = sum(1 for p in P if p['zone'] in ('ok', 'att', 'ko'))
print('produits', len(P), 'chiffrés', nOk, 'zones', {z: sum(1 for p in P if p['zone'] == z) for z in ZL}, 'groupes', len(groupes))

# ---------- Excel
wb = Workbook(); ws = wb.active; ws.title = 'Produits'
F_T = Font(name='Arial', size=13, bold=True, color='8D1D2C'); F_S = Font(name='Arial', size=9, color='7A736A'); F_H = Font(name='Arial', size=9, bold=True, color='FFFFFF'); F_B = Font(name='Arial', size=9)
FILL_H = PatternFill('solid', fgColor='221E1A'); thin = Side(style='thin', color='E6E0D8'); BD = Border(bottom=thin)
ws['A1'] = 'Atelier by – Halle · les produits et leur coefficient'; ws['A1'].font = F_T
ws['A2'] = f'Coefficient = prix de vente HT de la pièce ÷ coût de recette net du jour (API du panel, lu le {LU}). Zones : vert ≥ × 2,5 (marge brute ≥ 60 %), orange × 1,67 à 2,5 (marge 40 à 60 %), rouge < × 1,67 (marge < 40 %) ; objectif du réseau × 3,13 (matière 32 %). Cible : le coefficient attendu de la catégorie dans le panel. « Recette à vérifier » : coût au-dessus du prix ou sous 5 % de celui-ci.'; ws['A2'].font = F_S
ws['A2'].alignment = Alignment(wrap_text=True, vertical='top'); ws.merge_cells('A2:Q2'); ws.row_dimensions[2].height = 42
cols = ['Groupe', 'Catégorie', 'Produit', 'Prix TTC', 'TVA', 'Prix HT', 'Coût matière HT', 'Coefficient', 'Marge brute', 'Food cost', 'Zone', 'Cible catégorie', 'Écart à la cible', 'Prix TTC pour la cible', 'Prix TTC objectif réseau', 'Durée de vie', 'Webshop']
for i, c in enumerate(cols, 1):
    x = ws.cell(4, i, c); x.font = F_H; x.fill = FILL_H; x.alignment = Alignment(horizontal='left' if i <= 3 else 'right', vertical='center', wrap_text=True)
ws.row_dimensions[4].height = 30
FMT = {4: '0.00 €', 5: '0 %', 6: '0.00 €', 7: '0.00 €', 8: '× 0.00', 9: '0.0 %', 10: '0.0 %', 12: '× 0.00', 13: '+0.00;−0.00', 14: '0.00 €', 15: '0.00 €'}
r = 5
for p in P:
    fill = PatternFill('solid', fgColor=teinte(GC[p['groupe']], 0.86))
    vals = [p['groupe'], p['cat'], p['nom'], p['prixTtc'], (p['tva'] or 0) / 100 if p['tva'] is not None else None, p['prixHt'], p['cout'], p['coef'] if p['zone'] != 'mu' else None,
            (p['margePct'] / 100 if p['margePct'] is not None and p['zone'] != 'ab' else None), (p['foodPct'] / 100 if p['foodPct'] is not None and p['zone'] != 'ab' else None), ZL[p['zone']],
            p['cible'], (p['ecartCible'] if p['zone'] in ('ok', 'att', 'ko') else None), (p['prixCible'] if p['cout'] else None), (p['prixObjectif'] if p['cout'] else None), p['vie'] or '', 'oui' if p['webshop'] else '']
    for i, v in enumerate(vals, 1):
        x = ws.cell(r, i, v); x.font = F_B; x.border = BD; x.fill = fill
        if i in FMT: x.number_format = FMT[i]
        x.alignment = Alignment(horizontal='left' if i in (1, 2, 3, 11, 16, 17) else 'right')
    zc = ZC[p['zone']]; ws.cell(r, 8).fill = PatternFill('solid', fgColor=zc[1]); ws.cell(r, 8).font = Font(name='Arial', size=9, bold=True, color=zc[0]); ws.cell(r, 11).font = Font(name='Arial', size=8, color=zc[0])
    r += 1
for i, w in enumerate([18, 22, 34, 9, 6, 9, 12, 10, 10, 9, 26, 9, 9, 12, 12, 9, 8], 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'D5'; ws.auto_filter.ref = f'A4:Q{r-1}'
def feuille(titre, entetes, lignes, largeurs, fmts, sous=''):
    w = wb.create_sheet(titre); w['A1'] = titre; w['A1'].font = F_T
    if sous: w['A2'] = sous; w['A2'].font = F_S
    for i, c in enumerate(entetes, 1):
        x = w.cell(4, i, c); x.font = F_H; x.fill = FILL_H; x.alignment = Alignment(horizontal='left' if i <= 2 else 'right', vertical='center', wrap_text=True)
    w.row_dimensions[4].height = 30
    rr = 5
    for l in lignes:
        for i, v in enumerate(l, 1):
            x = w.cell(rr, i, v); x.font = F_B; x.border = BD; x.alignment = Alignment(horizontal='left' if i <= 2 else 'right')
            if i in fmts and v is not None: x.number_format = fmts[i]
        rr += 1
    for i, lw in enumerate(largeurs, 1): w.column_dimensions[get_column_letter(i)].width = lw
    w.freeze_panes = 'A5'
    return w
w2 = feuille('Par catégorie', ['Groupe', 'Catégorie', 'Produits', 'Chiffrés', 'Coefficient médian', 'Minimum', 'Maximum', 'Cible du panel', 'Vert', 'Orange', 'Rouge', 'Sans coût', 'À vérifier'],
    [[a['groupe'], a['nom'], a['n'], a['chiffres'], a['med'], a['mn'], a['mx'], a['cible'], a['z']['ok'], a['z']['att'], a['z']['ko'], a['z']['mu'], a['z']['ab']] for a in AC],
    [18, 24, 9, 9, 11, 9, 9, 10, 7, 7, 7, 9, 9], {5: '× 0.00', 6: '× 0.00', 7: '× 0.00', 8: '× 0.00'}, 'Médiane, minimum et maximum sur les produits chiffrés, anomalies exclues.')
for rr in range(5, 5 + len(AC)): w2.cell(rr, 1).fill = PatternFill('solid', fgColor=teinte(GC[w2.cell(rr, 1).value], 0.86))
w3 = feuille('Par groupe', ['Groupe', '', 'Produits', 'Chiffrés', 'Coefficient médian', 'Minimum', 'Maximum', 'Vert', 'Orange', 'Rouge', 'Sans coût', 'À vérifier'],
    [[a['nom'], '', a['n'], a['chiffres'], a['med'], a['mn'], a['mx'], a['z']['ok'], a['z']['att'], a['z']['ko'], a['z']['mu'], a['z']['ab']] for a in AG],
    [22, 3, 9, 9, 11, 9, 9, 7, 7, 7, 9, 9], {5: '× 0.00', 6: '× 0.00', 7: '× 0.00'})
for rr in range(5, 5 + len(AG)): w3.cell(rr, 1).fill = PatternFill('solid', fgColor=teinte(GC[w3.cell(rr, 1).value], 0.86))
av = [p for p in P if p['zone'] in ('ab', 'mu')]
def motif(p):
    if p['zone'] == 'ab': return 'coût de recette invraisemblable face au prix (' + ('au-dessus du prix' if p['coef'] <= 1 else 'sous 5 % du prix') + ')'
    if p['cout'] is None and p['prixHt'] is None: return 'ni coût de recette ni prix'
    return 'sans coût de recette' if p['cout'] is None else 'sans prix'
feuille('À vérifier', ['Groupe', 'Catégorie', 'Produit', 'Prix TTC', 'Coût matière HT', 'Coefficient', 'Motif'],
    [[p['groupe'], p['cat'], p['nom'], p['prixTtc'], p['cout'], p['coef'], motif(p)] for p in av], [18, 22, 34, 9, 12, 10, 60], {4: '0.00 €', 5: '0.00 €', 6: '× 0.00'},
    f'{len(av)} produits sans coefficient utilisable : sans coût de recette au panel, sans prix, ou recette invraisemblable.')
wb.save(SP + '/coef/Halle-produits-coefficients.xlsx'); print('xlsx écrit')

# ---------- HTML → PDF
Hs = lambda s: html.escape(str(s))
nf = lambda v, dd=2: '—' if v is None else f'{v:,.{dd}f}'.replace(',', ' ').replace('.', ',')
LOGO = 'data:image/png;base64,' + base64.b64encode(open('/home/user/consultant_BO/public/assets/img/logo.png', 'rb').read()).decode()
css = '''
@font-face{font-family:'Gotham';src:url('file:///home/user/consultant_BO/public/assets/ds/fonts/Gotham_Regular_New.otf');font-weight:400}
@font-face{font-family:'Gotham';src:url('file:///home/user/consultant_BO/public/assets/ds/fonts/Gotham_Medium.otf');font-weight:500}
@font-face{font-family:'Vank';src:url('file:///home/user/consultant_BO/public/assets/ds/fonts/GC_Vank.ttf')}
@page{size:A4;margin:12mm 12mm 14mm}
body{font-family:Gotham,Helvetica,Arial,sans-serif;color:#221E1A;font-size:8.2pt;line-height:1.35;margin:0}
.hd{display:flex;justify-content:space-between;align-items:center;border-bottom:2px solid #8D1D2C;padding-bottom:2mm;margin-bottom:4mm}
.hd img{height:6mm}.hd .dr{font-size:7.5pt;letter-spacing:.08em;text-transform:uppercase;color:#7a736a}
h1{font-family:Vank,Georgia,serif;font-weight:400;font-size:20pt;margin:0 0 1mm}
.s{font-size:7.6pt;color:#7a736a;line-height:1.45;margin:0 0 3mm}
.tuiles{display:grid;grid-template-columns:repeat(5,1fr);gap:2.5mm;margin:0 0 4mm}
.tuile{border:1px solid #e6e0d8;border-radius:6px;padding:2mm 2.5mm;background:#fbf9f5}.tuile .k{font-size:6.5pt;text-transform:uppercase;letter-spacing:.06em;color:#7a736a}.tuile .v{font-family:Vank,Georgia,serif;font-size:15pt;color:#8D1D2C;margin-top:.5mm}.tuile .t{font-size:6.8pt;color:#7a736a}
.leg{display:flex;gap:4mm;font-size:7pt;color:#5a534b;margin:0 0 4mm}.leg i{display:inline-block;width:9px;height:9px;border-radius:2px;vertical-align:-1px;margin-right:3px}
.grp{margin:0 0 4mm;page-break-inside:auto}
.gh{display:flex;justify-content:space-between;align-items:baseline;color:#fff;padding:1.4mm 2.5mm;border-radius:4px 4px 0 0;font-size:9pt;font-weight:500;page-break-after:avoid}
.gh small{font-weight:400;font-size:7pt;opacity:.9}
table{width:100%;border-collapse:collapse}thead{display:table-header-group}
th{font-size:6.3pt;letter-spacing:.05em;text-transform:uppercase;color:#7a736a;font-weight:500;text-align:right;padding:1mm 1.6mm;border-bottom:1pt solid #221E1A}
td{padding:.9mm 1.6mm;border-bottom:.5pt solid #EAE3D8;text-align:right;white-space:nowrap;font-variant-numeric:tabular-nums}
tr{page-break-inside:avoid}
th.l,td.l{text-align:left;white-space:normal}td.c{font-size:6.6pt;color:#7a736a}
.chip{display:inline-block;min-width:11mm;text-align:center;padding:.2mm 1.4mm;border-radius:3px;font-weight:500}
.ok{background:#D7EBD9;color:#2D7A3E}.att{background:#FDE6D0;color:#A85A10}.ko{background:#F6D3D6;color:#C0182B}.ab{background:#E6DDF2;color:#6A3FA0}.mu{background:#EEECE8;color:#7a736a}
.note{border-left:3px solid #8D1D2C;background:#fbf9f5;padding:1.5mm 2.5mm;font-size:7.4pt;color:#5a534b;margin:3mm 0}
.cats{columns:3;column-gap:5mm;font-size:7.2pt}.cats div{break-inside:avoid;padding:.4mm 0;border-bottom:.5pt solid #EAE3D8;display:flex;justify-content:space-between}
'''
parts = [f'<div class="hd"><img src="{LOGO}" alt=""><span class="dr">Atelier by – Halle · les coefficients produit par produit · lu le {LU}</span></div>']
parts.append(f'<h1>Les produits de Halle et leur coefficient</h1><div class="s">Coefficient = prix de vente HT de la pièce ÷ coût de recette net du jour, tels que le panel les calcule pour le magasin. Le groupe est la famille du panel, la catégorie sa sous-catégorie ; chaque groupe porte sa couleur. La cible est le coefficient attendu de la catégorie dans le panel ; le prix pour la cible est le prix TTC qui l’atteindrait, aux 5 centimes supérieurs.</div>')
Z = {z: sum(1 for p in P if p['zone'] == z) for z in ZL}
parts.append('<div class="tuiles">' + ''.join(f'<div class="tuile"><div class="k">{k}</div><div class="v">{v}</div><div class="t">{t}</div></div>' for k, v, t in [
    ('Produits disponibles', len(P), f'{len(groupes)} groupes, {len(AC)} catégories'), ('Avec un coefficient', nOk, f'médiane × {nf(med([p["coef"] for p in P if p["zone"] in ("ok","att","ko")]))}'),
    ('Dans le vert', Z['ok'], 'marge brute ≥ 60 %'), ('Sous × 1,67', Z['ko'], 'marge brute < 40 %'), ('À vérifier', Z['ab'] + Z['mu'], f'{Z["ab"]} recettes invraisemblables, {Z["mu"]} sans coût ou sans prix')]) + '</div>')
parts.append('<div class="leg">' + ''.join(f'<span><i class="{z}" style="background:#{ZC[z][1]};border:1px solid #{ZC[z][0]}"></i>{ZL[z]}</span>' for z in ZL) + f'<span>objectif du réseau × {nf(S["objectif"])}</span></div>')
for g in groupes:
    ps = [p for p in P if p['groupe'] == g]; a = [x for x in AG if x['nom'] == g][0]
    rows = ''
    for p in ps:
        if p['zone'] == 'mu': continue
        rows += (f'<tr><td class="l c">{Hs(p["cat"]) if p["cat"] != g else ""}</td><td class="l">{Hs(p["nom"])}</td><td>{nf(p["prixTtc"])} €</td><td>{nf(p["cout"])} €</td>'
                 f'<td><span class="chip {p["zone"]}">× {nf(p["coef"])}</span></td><td>{nf(p["margePct"],0) + " %" if p["zone"] != "ab" else "—"}</td><td>{("× " + nf(p["cible"])) if p["cible"] else "—"}</td>'
                 f'<td>{(nf(p["prixCible"]) + " €") if p["prixCible"] and p["zone"] != "ab" else "—"}</td><td>{(nf(p["prixObjectif"]) + " €") if p["prixObjectif"] and p["zone"] != "ab" else "—"}</td></tr>')
    sans = [p['nom'] for p in ps if p['zone'] == 'mu']
    parts.append(f'<div class="grp"><div class="gh" style="background:#{GC[g]}"><span>{Hs(g)}</span><small>{a["n"]} produits · {a["chiffres"]} chiffrés · médiane {("× " + nf(a["med"])) if a["med"] else "—"} · {a["z"]["ok"]} vert, {a["z"]["att"]} orange, {a["z"]["ko"]} rouge{(" · " + str(a["z"]["ab"]) + " à vérifier") if a["z"]["ab"] else ""}</small></div>'
                 + (f'<table><thead><tr><th class="l">Catégorie</th><th class="l">Produit</th><th>Prix TTC</th><th>Coût matière</th><th>Coefficient</th><th>Marge brute</th><th>Cible</th><th>Prix pour la cible</th><th>Prix objectif réseau</th></tr></thead><tbody>{rows}</tbody></table>' if rows else '')
                 + (f'<div class="s" style="margin:1.5mm 0 0;padding:0 2mm">Sans coût de recette ou sans prix ({len(sans)}) : {Hs(", ".join(sans))}.</div>' if sans else '') + '</div>')
parts.append('<div class="note">Les coefficients « à vérifier » viennent de recettes dont le coût dépasse le prix de vente ou n’en fait pas 5 % : une recette chiffrée pour un plateau ou une boîte quand le prix est à la pièce, ou l’inverse. Ils sont écartés des médianes. Les produits sans coût de recette (boissons, cafés, extras) n’ont pas de coefficient tant que le panel ne les chiffre pas.</div>')
doc = f'<!doctype html><html lang="fr"><head><meta charset="utf-8"><title>Halle · coefficients produits</title><style>{css}</style></head><body>{"".join(parts)}</body></html>'
open(SP + '/coef/liste.html', 'w', encoding='utf-8').write(doc); print('html écrit', len(doc))
